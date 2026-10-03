"""Local text extraction for Job Descriptions and CVs.

Nothing here calls Claude -- this module only turns PDF/DOCX/TXT files into
plain text so that text can be sent to Claude afterwards. No candidate
information is invented here: if a file has no readable text, the caller
gets None and must record that honestly rather than guessing content.
"""

from __future__ import annotations

import os
from typing import Optional

import pymupdf as fitz  # PyMuPDF (the `fitz` import name is deprecated upstream)
import docx  # python-docx

NO_TEXT_MARKER = "Could not extract readable CV text"


def extract_pdf_text(path: str) -> str:
    """Extract text from a PDF using PyMuPDF."""
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def extract_docx_text(path: str) -> str:
    """Extract text (paragraphs + table cells) from a DOCX using python-docx."""
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def extract_txt_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def extract_document_text(path: str) -> Optional[str]:
    """Dispatch by file extension.

    Returns the extracted text, or None if the file could not be read or
    contains no readable text at all. Callers must treat None as
    "Could not extract readable CV text" -- never substitute guessed
    content.
    """
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            text = extract_pdf_text(path)
        elif ext == ".docx":
            text = extract_docx_text(path)
        elif ext == ".txt":
            text = extract_txt_text(path)
        else:
            return None
    except Exception:
        # Corrupt file, password-protected, unsupported internal format, etc.
        return None

    if not text or not text.strip():
        return None
    return text
